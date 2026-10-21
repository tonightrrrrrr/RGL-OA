"""把 docs 下的 Markdown 文档转成排版好的 Word。
用法：python md2docx.py <输入.md> <输出.docx> <页脚文字>
依赖：python-docx==1.1.2
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CJK, MONO = "Microsoft YaHei", "Consolas"
INK, HEAD = (28, 38, 43), (0x1F, 0x33, 0x3B)
ACCENT, STRIPE, CODE_BG = "304954", "F2F6F7", "F4F6F8"
TEXT_WIDTH_CM = 16.3

# OOXML 要求子元素按固定顺序排列，这里列出必须排在 shd / tblBorders 等元素之后的兄弟元素
PPR_AFTER_SHD = ["tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
                 "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing",
                 "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap",
                 "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
TCPR_AFTER_SHD = ["noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]
TBLPR_AFTER_BORDERS = ["shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
TBLPR_AFTER_CELLMAR = ["tblLook", "tblCaption", "tblDescription"]


def insert_ordered(parent, child, successors):
    names = {qn(f"w:{n}") for n in successors}
    for existing in parent:
        if existing.tag in names:
            existing.addprevious(child)
            return
    parent.append(child)


def shade(pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    insert_ordered(pr, shd, PPR_AFTER_SHD if pr.tag == qn("w:pPr") else TCPR_AFTER_SHD)


def set_font(run, name=CJK, size=None, bold=None, color=None, east=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), east or (CJK if name == MONO else name))
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_inline(p, text, size=10.5, color=INK, bold=False):
    """行内 `代码`、**加粗**、[链接](地址)（只保留文字）"""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            set_font(p.add_run(part[1:-1]), MONO, size - 0.8, bold, (34, 67, 78))
        elif part.startswith("**") and part.endswith("**"):
            set_font(p.add_run(part[2:-2]), CJK, size, True, color)
        else:
            set_font(p.add_run(part), CJK, size, bold, color)


def visual_len(s):
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in re.sub(r"[`*]", "", s))


def add_table(doc, lines):
    rows = [[c.strip() for c in l.strip().strip("|").split("|")] for l in lines]
    rows = [r for r in rows if not all(re.fullmatch(r":?-{3,}:?", c) for c in r)]
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    rows = [r + [""] * (ncol - len(r)) for r in rows]
    lens = [max(min(visual_len(r[c]), 60) for r in rows) + 4 for c in range(ncol)]
    ratios = [min(max(l / sum(lens), 0.08), 0.55) for l in lens]
    widths = [Cm(TEXT_WIDTH_CM * r / sum(ratios)) for r in ratios]
    size = 9 if ncol <= 3 else 8.5 if ncol <= 4 else 8
    table = doc.add_table(rows=0, cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{side}")
        for k, v in (("val", "single"), ("sz", "4"), ("color", "C9D1D4")):
            b.set(qn(f"w:{k}"), v)
        borders.append(b)
    insert_ordered(tbl_pr, borders, TBLPR_AFTER_BORDERS)
    margins = OxmlElement("w:tblCellMar")
    for side, w in (("top", 50), ("left", 90), ("bottom", 50), ("right", 90)):
        m = OxmlElement(f"w:{side}")
        m.set(qn("w:w"), str(w))
        m.set(qn("w:type"), "dxa")
        margins.append(m)
    insert_ordered(tbl_pr, margins, TBLPR_AFTER_CELLMAR)
    for ri, data in enumerate(rows):
        row = table.add_row()
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))  # 跨页重复表头
        for ci, value in enumerate(data):
            cell = row.cells[ci]
            cell.width = widths[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            if ri == 0:
                add_inline(p, value, size, (255, 255, 255), True)
                shade(cell._tc.get_or_add_tcPr(), ACCENT)
            else:
                add_inline(p, value, size)
                if ri % 2 == 0:
                    shade(cell._tc.get_or_add_tcPr(), STRIPE)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_code(doc, lines):
    widest = max((visual_len(l) for l in lines), default=0)
    size = 8.5 if widest <= 92 else 7.5 if widest <= 104 else 6.5
    for i, line in enumerate(lines):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.space_before = Pt(4) if i == 0 else Pt(0)
        pf.space_after = Pt(4) if i == len(lines) - 1 else Pt(0)
        pf.line_spacing = 1.0
        pf.left_indent = Cm(0.2)
        pf.keep_with_next = i < len(lines) - 1
        shade(p._p.get_or_add_pPr(), CODE_BG)
        # 中文字号放大 1.1 倍，使一个中文正好占两个英文字宽，图示边框才能对齐
        for part in re.findall(r"[⺀-￿]+|[^⺀-￿]+", line) or [" "]:
            wide = ord(part[0]) > 0x2E80
            set_font(p.add_run(part), "NSimSun" if wide else MONO, round(size * 1.1, 1) if wide else size,
                     color=(40, 50, 56), east="NSimSun")


def add_field(paragraph, instr, size=8, color=(100, 110, 115), placeholder=None):
    for kind, text in (("begin", None), (None, instr), ("separate", None), ("text", placeholder), ("end", None)):
        if kind == "text":
            if placeholder:
                set_font(paragraph.add_run(placeholder), CJK, size, color=color)
            continue
        run = paragraph.add_run()
        set_font(run, CJK, size, color=color)
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            run._r.append(it)


def setup(doc, footer_text):
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21)
    sec.top_margin, sec.bottom_margin = Cm(2.3), Cm(2.2)
    sec.left_margin, sec.right_margin = Cm(2.35), Cm(2.35)
    for name, size, before, after in [("Normal", 10.5, 0, 6), ("Title", 22, 0, 12), ("Heading 1", 16, 6, 10),
                                      ("Heading 2", 13, 12, 6), ("Heading 3", 11.5, 10, 4)]:
        st = doc.styles[name]
        st.font.name = CJK
        st.font.size = Pt(size)
        st.font.bold = name != "Normal"
        st.font.color.rgb = RGBColor(*(INK if name == "Normal" else HEAD))
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK)
        pf = st.paragraph_format
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.line_spacing = 1.3 if name == "Normal" else 1.15
        pf.keep_with_next = name != "Normal"
    for name in ("List Bullet", "List Bullet 2"):
        doc.styles[name].font.name = CJK
        doc.styles[name].font.size = Pt(10.5)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run(footer_text + "　　第 "), CJK, 8, color=(100, 110, 115))
    add_field(footer, "PAGE")
    set_font(footer.add_run(" 页"), CJK, 8, color=(100, 110, 115))


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def main(src, out, footer_text):
    doc = Document()
    setup(doc, footer_text)
    lines = Path(src).read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        raw, line = lines[i], lines[i].strip()
        if not line or line == "---" or line.startswith("<!--"):
            i += 1
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            add_code(doc, code)
            i += 1
            continue
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(120)
            add_inline(p, line[2:], 22, HEAD, True)
        elif line.startswith("## "):
            title = line[3:]
            page_break(doc)
            if title == "目录":
                # 手写目录换成 Word 自动目录；目录标题不用标题样式，避免出现在目录里
                add_inline(doc.add_paragraph(), title, 16, HEAD, True)
                add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', 10.5, INK, "（目录将在 Word 中更新）")
                i += 1
                while i < len(lines) and (not lines[i].strip() or re.match(r"^(\d+\.|附录)", lines[i].strip())):
                    i += 1
                continue
            add_inline(doc.add_paragraph(style="Heading 1"), title, 16, HEAD, True)
        elif line.startswith("### "):
            add_inline(doc.add_paragraph(style="Heading 2"), line[4:], 13, HEAD, True)
        elif line.startswith("#### "):
            add_inline(doc.add_paragraph(style="Heading 3"), line[5:], 11.5, HEAD, True)
        elif re.match(r"^-\s", line):
            p = doc.add_paragraph(style="List Bullet 2" if raw.startswith("  ") else "List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line[2:])
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(0.75), Cm(-0.55)
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, line)
        else:
            add_inline(doc.add_paragraph(), line)
        i += 1
    doc.save(out)
    print(out)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else "企业OA文件智能归档智能体")